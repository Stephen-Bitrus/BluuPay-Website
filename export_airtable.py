#!/usr/bin/env python3
"""Export the public Airtable shared view to Excel."""

from __future__ import annotations

import json
import re
from collections import Counter
from datetime import datetime, timezone
from urllib.parse import parse_qs, urlencode, urlparse, urlunparse

import requests
from openpyxl import Workbook
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter

SHARE_URL = (
    "https://airtable.com/appcim1kbCvNaNrae/shrY0E6dt9elq5j5k/"
    "tblYUkt345u2u7giJ/viwSUqTO4Uvmg8cFl"
)
APP_ID = "appcim1kbCvNaNrae"
OUTPUT = "155_March_Opportunities_African_Startups.xlsx"
USER_AGENT = (
    "Mozilla/5.0 (Windows NT 10.0; Win64; x64) "
    "AppleWebKit/537.36 (KHTML, like Gecko) Chrome/120.0.0.0 Safari/537.36"
)


def fetch_view_payload() -> dict:
    session = requests.Session()
    page = session.get(
        SHARE_URL,
        headers={"User-Agent": USER_AGENT, "Accept": "text/html"},
        timeout=60,
    )
    page.raise_for_status()
    match = re.search(r'urlWithParams:\s*"((?:\\.|[^"\\])*)"', page.text)
    if not match:
        raise RuntimeError("Could not find Airtable data URL in the share page.")
    path = match.group(1).encode().decode("unicode_escape")
    parsed = urlparse("https://airtable.com" + path)
    query = parse_qs(parsed.query)
    params = json.loads(query["stringifiedObjectParams"][0])
    params["allowMsgpackOfResult"] = False
    query["stringifiedObjectParams"] = [json.dumps(params, separators=(",", ":"))]
    api_url = urlunparse(
        (
            parsed.scheme,
            parsed.netloc,
            parsed.path,
            parsed.params,
            urlencode({key: values[0] for key, values in query.items()}),
            parsed.fragment,
        )
    )
    response = session.get(
        api_url,
        headers={
            "User-Agent": USER_AGENT,
            "Accept": "application/json",
            "X-Requested-With": "XMLHttpRequest",
            "x-airtable-application-id": APP_ID,
            "x-airtable-inter-service-client": "webClient",
            "x-time-zone": "UTC",
            "Referer": SHARE_URL,
        },
        timeout=120,
    )
    response.raise_for_status()
    payload = response.json()
    if payload.get("msg") != "SUCCESS":
        raise RuntimeError(f"Airtable read failed: {payload.get('msg')}")
    return payload["data"]


def cell_value(column: dict, raw) -> str:
    if raw is None:
        return ""
    if column["type"] == "select":
        choice = (column.get("typeOptions") or {}).get("choices", {}).get(raw)
        return choice["name"] if choice else str(raw)
    if isinstance(raw, list):
        parts = []
        for item in raw:
            if isinstance(item, dict):
                parts.append(
                    item.get("name")
                    or item.get("url")
                    or item.get("foreignRowDisplayName")
                    or str(item)
                )
            else:
                parts.append(str(item))
        return ", ".join(parts)
    if isinstance(raw, dict):
        return raw.get("name") or raw.get("url") or json.dumps(raw)
    return str(raw)


def write_workbook(data: dict, path: str) -> tuple[int, int]:
    schema = data["tableSchemas"][0]
    table_data = data["tableDatas"][0]
    view = table_data["viewDatas"][0]
    columns_by_id = {column["id"]: column for column in schema["columns"]}
    visible_columns = [
        columns_by_id[item["columnId"]]
        for item in view["columnOrder"]
        if item.get("visibility", True)
    ]
    rows_by_id = {row["id"]: row for row in table_data["rows"]}
    ordered_rows = [
        rows_by_id[item["rowId"]]
        for item in view["rowOrder"]
        if item.get("visibility", True) and item["rowId"] in rows_by_id
    ]

    workbook = Workbook()
    sheet = workbook.active
    sheet.title = schema["name"][:31] or "Opportunities"

    headers = [column["name"] for column in visible_columns]
    header_fill = PatternFill("solid", fgColor="1B4332")
    header_font = Font(name="Calibri", bold=True, color="FFFFFF", size=11)
    body_font = Font(name="Calibri", size=11)
    wrap = Alignment(wrap_text=True, vertical="top")
    thin = Border(
        left=Side(style="thin", color="D8E2DC"),
        right=Side(style="thin", color="D8E2DC"),
        top=Side(style="thin", color="D8E2DC"),
        bottom=Side(style="thin", color="D8E2DC"),
    )
    alt_fill = PatternFill("solid", fgColor="F4FBF7")
    open_fill = PatternFill("solid", fgColor="D8F3DC")
    closed_fill = PatternFill("solid", fgColor="FADDD8")
    link_font = Font(name="Calibri", size=11, color="0B6E4F", underline="single")

    sheet.append(headers)
    for cell in sheet[1]:
        cell.fill = header_fill
        cell.font = header_font
        cell.alignment = Alignment(wrap_text=True, vertical="center")

    status_col = headers.index("Status") + 1 if "Status" in headers else None
    link_col = headers.index("Application Link") + 1 if "Application Link" in headers else None
    url_re = re.compile(r"^https?://", re.I)

    for row_index, row in enumerate(ordered_rows, start=2):
        values = [
            cell_value(column, row.get("cellValuesByColumnId", {}).get(column["id"]))
            for column in visible_columns
        ]
        sheet.append(values)
        for col_index, cell in enumerate(sheet[row_index], start=1):
            cell.font = body_font
            cell.alignment = wrap
            cell.border = thin
            if row_index % 2 == 0:
                cell.fill = alt_fill
            if col_index == link_col:
                value = str(cell.value or "").strip()
                if url_re.match(value):
                    cell.hyperlink = value
                    cell.font = link_font
            if col_index == status_col:
                if cell.value == "Open":
                    cell.fill = open_fill
                elif cell.value == "Closed":
                    cell.fill = closed_fill

    last_col = get_column_letter(len(headers))
    sheet.auto_filter.ref = f"A1:{last_col}{sheet.max_row}"
    sheet.freeze_panes = "B2"
    sheet.sheet_view.showGridLines = False
    sheet.print_title_rows = "1:1"
    sheet.page_setup.orientation = "landscape"
    sheet.page_setup.fitToPage = True
    sheet.page_setup.fitToWidth = 1

    widths = {
        "Program Name": 42,
        "Organizing Entity": 28,
        "Target: Country/Region": 28,
        "Target: Sector": 32,
        "Target: Stage": 28,
        "Target: Who": 36,
        "Benefit: Funding": 36,
        "Benefit: Additional": 42,
        "Brief Description": 48,
        "Application Link": 36,
        "Deadline": 16,
        "Status": 12,
    }
    for index, name in enumerate(headers, start=1):
        sheet.column_dimensions[get_column_letter(index)].width = widths.get(name, 22)
    sheet.row_dimensions[1].height = 22
    for row_index in range(2, sheet.max_row + 1):
        sheet.row_dimensions[row_index].height = 48

    status_counts: Counter[str] = Counter()
    deadline_counts: Counter[str] = Counter()
    status_column = next(column for column in visible_columns if column["name"] == "Status")
    deadline_column = next(column for column in visible_columns if column["name"] == "Deadline")
    for row in ordered_rows:
        cells = row.get("cellValuesByColumnId", {})
        status_counts[cell_value(status_column, cells.get(status_column["id"]))] += 1
        deadline_counts[cell_value(deadline_column, cells.get(deadline_column["id"]))] += 1

    summary = workbook.create_sheet("Summary", 1)
    summary["A1"] = "Export snapshot"
    summary["A1"].font = Font(name="Calibri", bold=True, size=14, color="1B4332")
    summary["A3"] = "Status"
    summary["B3"] = "Count"
    summary["A3"].font = header_font
    summary["B3"].font = header_font
    summary["A3"].fill = header_fill
    summary["B3"].fill = header_fill
    row_num = 4
    for label, count in status_counts.most_common():
        summary[f"A{row_num}"] = label
        summary[f"B{row_num}"] = count
        row_num += 1
    row_num += 1
    summary[f"A{row_num}"] = "Deadline"
    summary[f"B{row_num}"] = "Count"
    summary[f"A{row_num}"].font = header_font
    summary[f"B{row_num}"].font = header_font
    summary[f"A{row_num}"].fill = header_fill
    summary[f"B{row_num}"].fill = header_fill
    row_num += 1
    for label, count in deadline_counts.most_common():
        summary[f"A{row_num}"] = label
        summary[f"B{row_num}"] = count
        row_num += 1
    summary.column_dimensions["A"].width = 28
    summary.column_dimensions["B"].width = 12

    source = workbook.create_sheet("Source")
    source["A1"] = "Airtable export"
    source["A1"].font = Font(name="Calibri", bold=True, size=14, color="1B4332")
    meta = [
        ("Title", "155+ March Opportunities Built for African Startups"),
        ("Table", schema["name"]),
        ("View", view["id"]),
        ("Share URL", SHARE_URL),
        ("Records exported", len(ordered_rows)),
        ("Columns exported", len(visible_columns)),
        ("Exported at (UTC)", datetime.now(timezone.utc).strftime("%Y-%m-%d %H:%M:%S")),
        (
            "Notes",
            "All visible columns and rows from this shared grid view. "
            "Select fields are written as labels. Valid application URLs are clickable.",
        ),
    ]
    for index, (key, value) in enumerate(meta, start=3):
        source[f"A{index}"] = key
        source[f"B{index}"] = value
        source[f"A{index}"].font = Font(name="Calibri", bold=True)
        source[f"B{index}"].font = Font(name="Calibri")
        source[f"B{index}"].alignment = Alignment(wrap_text=True)
    source.column_dimensions["A"].width = 22
    source.column_dimensions["B"].width = 110
    source["B6"].hyperlink = SHARE_URL
    source["B6"].font = link_font

    workbook.save(path)
    return len(ordered_rows), len(visible_columns)


def main() -> None:
    data = fetch_view_payload()
    rows, cols = write_workbook(data, OUTPUT)
    print(f"Wrote {OUTPUT} ({rows} rows, {cols} columns)")


if __name__ == "__main__":
    main()
